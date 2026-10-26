import io
import json
import unittest
from openpyxl import load_workbook
from report_export import build_report_workbook


class ReportWorkbook(unittest.TestCase):
    def test_brand_is_appended_without_shifting_the_original_eighteen_columns(self):
        row={'category':'开发','fields':{'projectName':'合成品牌项目','brand':'=合成品牌'}}
        book=load_workbook(io.BytesIO(build_report_workbook([row])))
        self.assertEqual(book['开发'].max_column,19)
        self.assertEqual(book['开发']['R2'].value,'备注')
        self.assertEqual(book['开发']['S2'].value,'品牌')
        self.assertEqual(book['开发']['S3'].value,'=合成品牌')
        self.assertEqual(book['开发']['S3'].data_type,'s')
        self.assertEqual(book['重点项目'].max_column,19)
        from import_parser import read_import
        parsed=read_import(build_report_workbook([row]),'synthetic-brand.xlsx')
        raw=json.loads(parsed['sheets'][0]['rows'][0]['project']['reportData'])
        self.assertEqual(raw['fields']['品牌'],'=合成品牌')
    def test_four_sheets_match_eighteen_fields_and_priority_duplicates_only_as_view(self):
        row={'category':'开发','important':True,'currency':'CNY','quantityUnit':'件','priceTerms':'含税每件',
             'fields':{'customerName':'合成客户','projectName':'合成项目','quotePrice':'0.335','estimatedAnnualRevenue':'1234567890123.123456','statusNote':'等待测试','owner':'=1+2'}}
        book=load_workbook(io.BytesIO(build_report_workbook([row])))
        self.assertEqual(book.sheetnames,['开发','成交','EOL','重点项目'])
        self.assertEqual(book['开发'].max_column,18)
        self.assertEqual(book['开发']['E3'].value,'合成客户')
        self.assertEqual(book['开发']['G3'].value,'合成项目')
        self.assertEqual(book['开发']['K3'].value,0.335)
        self.assertEqual(book['开发']['M3'].value,'1234567890123.123456')
        self.assertEqual(book['开发']['A3'].data_type,'s')
        self.assertEqual(book['开发']['A3'].value,'=1+2')
        self.assertIn('CNY',book['开发']['K3'].comment.text)
        self.assertEqual(book['成交'].max_row,2)
        self.assertEqual(book['重点项目']['G3'].value,'合成项目')
        self.assertEqual(book['开发'].freeze_panes,'H3')

    def test_invalid_category_or_cell_too_long_refuses_incomplete_export(self):
        for row in [{'category':'未知','fields':{}},{'category':'开发','fields':{'notes':'a'*32768}}]:
            with self.assertRaises(ValueError):build_report_workbook([row])

    def test_legacy_high_scale_decimals_keep_the_original_visible_precision(self):
        rows=[{'category':'开发','fields':{'quotePrice':'0.1234567'}},{'category':'开发','fields':{'quotePrice':'0.0000001'}}]
        sheet=load_workbook(io.BytesIO(build_report_workbook(rows)))['开发']
        self.assertEqual(sheet['K3'].value,'0.1234567')
        self.assertEqual(sheet['K4'].value,'0.0000001')
        self.assertEqual(sheet['K4'].data_type,'s')
