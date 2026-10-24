"""Synthetic fixtures only: importing a file must never write project data."""
import csv
import datetime
import io
import json
import struct
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
try:
    import import_parser as parser
except ModuleNotFoundError as exc:
    if exc.name != 'import_parser':
        raise
    parser = None

BASIC = ['终端客户', 'ODM集成商', '模组厂', '项目名称', '应用', '目前阶段',
         '测试阶段需求数量', '项目预计总需求量', '预计MP日期']
REPORT = ['负责人', '日期', '产品线', '案子状况', '客户', '终端客户', '项目名称',
          '应用', '材料型号', '成品规格', '单价（含税）元', '预估年用量',
          '预估年营业额', '项目状态', '预估量产时间', '竞争对手', '竞争对手单价', '备注']


def csv_bytes(rows, encoding='utf-8-sig'):
    stream = io.StringIO(newline='')
    csv.writer(stream).writerows(rows)
    return stream.getvalue().encode(encoding)


def report_row(**values):
    row = [''] * 18
    row[6] = '合成测试项目'
    for key, value in values.items():
        row[REPORT.index(key)] = value
    return row


def xlsx_bytes(sheets):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def biff2_bytes(rows):
    # A real BIFF2 .xls stream, without an OLE wrapper.
    def record(kind, payload=b''):
        return struct.pack('<HH', kind, len(payload)) + payload
    data = record(0x0009, struct.pack('<HH', 2, 0x10)) + record(0x0042, struct.pack('<H', 936))
    for row_number, row in enumerate(rows):
        for column, value in enumerate(row):
            if value == '':
                continue
            prefix = struct.pack('<HH', row_number, column) + b'\0\0\0'
            if isinstance(value, (float, int)):
                data += record(0x0003, prefix + struct.pack('<d', value))
            else:
                encoded = str(value).encode('gbk')
                data += record(0x0004, prefix + bytes([len(encoded)]) + encoded)
    return data + record(0x000A)


class ImportParserTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(parser, '导入解析模块尚未实现')

    def first(self, content, filename='test.csv'):
        return parser.read_import(content, filename)['sheets'][0]

    def test_basic_csv_normalizes_numbers_and_keeps_multiline_csv(self):
        sheet = self.first(csv_bytes([BASIC, ['终端', '集成', '模组', '项目一', '应用,\n说明',
                                              'PVT', '0', '42', '2026-12-03']]))
        self.assertEqual(sheet['format'], 'basic')
        self.assertEqual(sheet['rows'][0]['rowNumber'], 2)
        self.assertEqual(sheet['rows'][0]['errors'], [])
        project = sheet['rows'][0]['project']
        self.assertEqual(project['estimatedTotalDemand'], 42)
        self.assertEqual(project['testDemandQuantity'], 0)
        self.assertEqual(project['application'], '应用,\n说明')
        self.assertEqual(project['expectedMpDate'], '2026-12-03')

    def test_gb18030_and_reordered_headers(self):
        rows = [list(reversed(BASIC)), list(reversed(['', '', '', '中文测试', '', '', '', '', '']))]
        self.assertEqual(self.first(csv_bytes(rows, 'gb18030'))['rows'][0]['project']['projectName'], '中文测试')

    def test_report_keeps_all_fields_without_guessing_old_fields(self):
        values = report_row(**{'负责人': '测试负责人', '案子状况': 2, '客户': '合成交易客户',
                               '终端客户': '合成终端', '应用': '测试应用', '材料型号': 'M-1',
                               '成品规格': '10*20', '单价（含税）元': '2.50元/件',
                               '预估年用量': '1万件', '预估量产时间': '2026年Q4'})
        sheet = self.first(csv_bytes([['合成报告标题'], REPORT, values]))
        self.assertEqual(sheet['format'], 'report')
        row = sheet['rows'][0]
        self.assertEqual(row['rowNumber'], 3)
        self.assertEqual(row['errors'], [])
        self.assertTrue(row['warnings'])
        project = row['project']
        self.assertEqual(project['terminalCustomer'], '合成终端')
        self.assertEqual(project['application'], '测试应用')
        for field in ['odmIntegrator', 'moduleFactory', 'stage']:
            self.assertFalse(project[field])
        for field in ['testDemandQuantity', 'estimatedTotalDemand', 'expectedMpDate']:
            self.assertIsNone(project[field])
        metadata = json.loads(project['reportData'])
        self.assertEqual(len(metadata['fields']), 18)
        self.assertEqual(metadata['fields']['单价（含税）元'], '2.50元/件')
        self.assertEqual(metadata['fields']['预估量产时间'], '2026年Q4')
        self.assertEqual(metadata['row'], 3)
        self.assertEqual(metadata['file'], 'test.csv')

    def test_plain_price_header_is_not_relabelled_as_tax_inclusive(self):
        header = REPORT.copy()
        header[10] = '单价'
        fields = json.loads(self.first(csv_bytes([header, report_row()]))['rows'][0]['project']['reportData'])['fields']
        self.assertIn('单价', fields)
        self.assertNotIn('单价（含税）元', fields)

    def test_header_whitespace_newlines_and_parentheses_are_normalized(self):
        header = REPORT.copy()
        header[0] = '负 责\n人'
        header[10] = '单价 (含税) 元'
        sheet = self.first(csv_bytes([header, report_row()]))
        self.assertEqual(sheet['format'], 'report')
        self.assertEqual(sheet['headers'][0], '负责人')

    def test_missing_and_slash_project_names_stay_blank_with_errors(self):
        sheet = self.first(csv_bytes([REPORT, report_row(**{'项目名称': '', '客户': '合成客户'}),
                                           report_row(**{'项目名称': '/'})]))
        for row in sheet['rows']:
            self.assertEqual(row['project']['projectName'], '')
            self.assertTrue(any(message.startswith('项目名称') for message in row['errors']))

    def test_basic_metadata_column_round_trips_raw_report(self):
        metadata = json.dumps({'fields': {'预估量产时间': '2026年Q4'}, 'sheet': '开发', 'file': 'old.xls', 'row': 3}, ensure_ascii=False)
        row = ['', '', '', '项目', '', '', '', '', '', metadata]
        project = self.first(csv_bytes([BASIC + ['开发报告原始资料'], row]))['rows'][0]['project']
        self.assertEqual(json.loads(project['reportData']), json.loads(metadata))

    def test_report_backup_json_keeps_original_fields_and_maps_current_columns(self):
        metadata = json.dumps({'fields': {'单价（含税）元': '2.50', '项目名称': '原项目',
                                         '预估量产时间': '2026年Q4'},
                               'sheet': '开发', 'file': 'source.xls', 'row': 7}, ensure_ascii=False)
        headers = REPORT.copy()
        headers[10] = '单价'
        headers += ['原单价表头', '开发报告原始资料', '来源工作表', '来源文件', '来源行号']
        values = report_row(**{'项目名称': '当前项目', '终端客户': '当前终端',
                               '应用': '当前应用', '预估量产时间': '2026-12-31'})
        values += ['单价（含税）元', metadata, '开发', 'source.xls', '7']
        sheet = self.first(csv_bytes([headers, values]))
        row = sheet['rows'][0]
        self.assertEqual(row['errors'], [])
        self.assertEqual(row['project']['projectName'], '当前项目')
        self.assertEqual(row['project']['terminalCustomer'], '当前终端')
        self.assertEqual(row['project']['application'], '当前应用')
        self.assertEqual(row['project']['expectedMpDate'], '2026-12-31')
        self.assertEqual(row['project']['reportData'], metadata)
        self.assertFalse(any('未映射' in warning or '不会导入' in warning for warning in sheet['warnings']))

    def test_repeated_report_preview_does_not_nest_or_grow_backup_metadata(self):
        metadata = json.dumps({'fields': {'单价（含税）元': '1.5', '项目名称': '原项目'},
                               'sheet': '开发', 'file': 'first.xls', 'row': 3}, ensure_ascii=False)
        headers = REPORT.copy()
        headers[10] = '单价'
        headers.append('开发报告原始资料')
        original_length = len(metadata)
        for _ in range(4):
            row = self.first(csv_bytes([headers, report_row() + [metadata]]))['rows'][0]
            metadata = row['project']['reportData']
            self.assertEqual(len(metadata), original_length)
            self.assertNotIn('开发报告原始资料', json.loads(metadata)['fields'])
            self.assertEqual(json.loads(metadata)['file'], 'first.xls')

    def test_bad_report_backup_is_a_row_error_and_not_silently_replaced(self):
        headers = REPORT + ['开发报告原始资料']
        row = self.first(csv_bytes([headers, report_row() + ['invalid json']]))['rows'][0]
        self.assertTrue(any(error.startswith('开发报告原始资料') for error in row['errors']))
        self.assertEqual(row['project']['reportData'], 'invalid json')

    def test_empty_report_backup_is_excluded_from_new_raw_fields(self):
        row = self.first(csv_bytes([REPORT + ['开发报告原始资料'], report_row() + ['']]))['rows'][0]
        self.assertEqual(row['errors'], [])
        fields = json.loads(row['project']['reportData'])['fields']
        self.assertEqual(len(fields), 18)
        self.assertNotIn('开发报告原始资料', fields)

    def test_public_header_canonicalization_is_available_to_report_export(self):
        canonicalize = getattr(parser, 'canonical_header', None)
        self.assertTrue(callable(canonicalize), '导出所需的公共表头规范化API尚未提供')
        self.assertEqual(canonicalize(' 单价（含税） 元\n'), '单价')
        self.assertEqual(canonicalize('负 责\n人'), '负责人')
        self.assertEqual(canonicalize('其他字段'), '其他字段')

    def test_invalid_rows_do_not_hide_following_valid_rows(self):
        sheet = self.first(csv_bytes([BASIC, ['', '', '', '', '', 'MP', '-1', '1.5', '2026-02-30'],
                                           ['', '', '', '有效项目', '', 'EVT', '1', '', '']]))
        self.assertGreaterEqual(len(sheet['rows'][0]['errors']), 4)
        self.assertEqual(sheet['rows'][1]['errors'], [])

    def test_empty_recognized_sheets_return_zero_rows(self):
        self.assertEqual(self.first(csv_bytes([REPORT]))['rows'], [])
        workbook = parser.read_import(xlsx_bytes([('开发', [REPORT]), ('成交', [REPORT])]), 'empty.xlsx')
        self.assertEqual([sheet['rows'] for sheet in workbook['sheets']], [[], []])

    def test_xlsx_preserves_multiple_sheet_sources_without_auto_merging(self):
        rows = [REPORT, report_row(**{'预估量产时间': datetime.date(2026, 11, 15)})]
        workbook = parser.read_import(xlsx_bytes([('开发', rows), ('重点项目', rows), ('说明', [['说明文字']])]), 'demo.xlsx')
        self.assertEqual([sheet['format'] for sheet in workbook['sheets']], ['report', 'report', 'unsupported'])
        self.assertEqual(workbook['sheets'][0]['rows'][0]['project']['expectedMpDate'], '2026-11-15')
        self.assertEqual(json.loads(workbook['sheets'][1]['rows'][0]['project']['reportData'])['sheet'], '重点项目')

    def test_xlsx_month_date_format_does_not_invent_day_precision(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(REPORT)
        sheet.append(report_row(**{'预估量产时间': datetime.date(2026, 11, 1)}))
        sheet['O2'].number_format = 'yyyy-mm'
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        row = self.first(stream.getvalue(), 'month.xlsx')['rows'][0]
        self.assertIsNone(row['project']['expectedMpDate'])
        self.assertEqual(json.loads(row['project']['reportData'])['fields']['预估量产时间'], '2026-11')

    def test_xls_binary_is_read_by_xlrd_without_assuming_csv(self):
        sheet = self.first(biff2_bytes([['unsupported heading'], ['value', 7]]), 'legacy.xls')
        self.assertEqual(sheet['format'], 'unsupported')

    def test_xls_chinese_report_rows_keep_values_and_quarter_precision(self):
        data = biff2_bytes([REPORT, report_row(**{'案子状况': 2, '预估量产时间': '2026年Q4'})])
        sheet = self.first(data, 'legacy.xls')
        self.assertEqual(sheet['format'], 'report')
        self.assertEqual(len(sheet['rows']), 1)
        project = sheet['rows'][0]['project']
        self.assertEqual(project['projectName'], '合成测试项目')
        self.assertIsNone(project['expectedMpDate'])
        self.assertEqual(json.loads(project['reportData'])['fields']['案子状况'], '2')

    def test_xlsx_merged_project_name_is_not_forward_filled(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(REPORT)
        sheet.append(report_row(**{'材料型号': 'SYNTH-1'}))
        sheet.append(report_row(**{'项目名称': '', '材料型号': 'SYNTH-2'}))
        sheet.merge_cells('G2:G3')
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        rows = self.first(stream.getvalue(), 'merged.xlsx')['rows']
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1]['project']['projectName'], '')
        self.assertTrue(rows[1]['errors'])

    def test_unknown_duplicate_recognized_headers_are_rejected(self):
        with self.assertRaisesRegex(ValueError, '重复'):
            self.first(csv_bytes([BASIC + ['项目名称'], ['', '', '', '项目', '', '', '', '', '', '另一个项目']]))

    def test_duplicate_extra_report_headers_cannot_silently_overwrite_raw_values(self):
        with self.assertRaisesRegex(ValueError, '重复'):
            self.first(csv_bytes([REPORT + ['附加字段', '附加字段'], report_row() + ['first', 'second']]))

    def test_unsupported_extension_and_corrupt_files_are_rejected(self):
        for filename, data in [('bad.txt', b'text'), ('bad.xls', b'not an excel file'),
                               ('bad.xlsx', b'not a zip'), ('bad.csv', b'\xff\xff\x00'), ('empty.csv', b'')]:
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                parser.read_import(data, filename)

    def test_malformed_csv_is_rejected(self):
        with self.assertRaises(ValueError):
            self.first(csv_bytes([BASIC]) + '"unterminated'.encode())

    def test_upload_size_limit(self):
        with self.assertRaisesRegex(ValueError, r'5\s*MB'):
            parser.read_import(b'x' * (5 * 1024 * 1024 + 1), 'large.csv')

    def test_xlsx_expansion_limit(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('oversized.xml', b'0' * (20 * 1024 * 1024 + 1))
        with self.assertRaisesRegex(ValueError, r'20\s*MB'):
            parser.read_import(stream.getvalue(), 'bomb.xlsx')

    def test_xlsx_xml_entities_are_rejected_before_loading(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<!DOCTYPE foo [<!ENTITY e "x">]><foo>&e;</foo>')
        with self.assertRaises(ValueError):
            parser.read_import(stream.getvalue(), 'entity.xlsx')

    def test_unknown_csv_header_is_a_file_error(self):
        with self.assertRaisesRegex(ValueError, '表头'):
            self.first(csv_bytes([['wrong header'], ['value']]))

    def test_xlsx_encrypted_zip_flag_is_rejected(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<Types/>')
        data = bytearray(stream.getvalue())
        local = data.index(b'PK\x03\x04')
        central = data.index(b'PK\x01\x02')
        data[local + 6] |= 1
        data[central + 8] |= 1
        with self.assertRaisesRegex(ValueError, '加密'):
            parser.read_import(bytes(data), 'encrypted.xlsx')

    def test_report_formula_is_preserved_without_execution(self):
        row = self.first(xlsx_bytes([('开发', [REPORT, report_row(**{'预估年营业额': '=1+2'})])]), 'formula.xlsx')['rows'][0]
        self.assertEqual(json.loads(row['project']['reportData'])['fields']['预估年营业额'], '=1+2')
        self.assertTrue(any('公式' in warning for warning in row['warnings']))

    def test_formula_project_name_requires_user_to_supply_confirmed_name(self):
        row = self.first(xlsx_bytes([('开发', [REPORT, report_row(**{'项目名称': '="合成项目"'})])]), 'formula.xlsx')['rows'][0]
        self.assertEqual(row['project']['projectName'], '')
        self.assertTrue(any(error.startswith('项目名称') for error in row['errors']))
        self.assertEqual(json.loads(row['project']['reportData'])['fields']['项目名称'], '="合成项目"')

    def test_basic_unknown_columns_warn_that_their_values_will_not_be_imported(self):
        sheet = self.first(csv_bytes([BASIC + ['其他资料'], ['', '', '', '项目', '', '', '', '', '', '额外值']]))
        self.assertTrue(any('不会导入' in warning for warning in sheet['warnings']))
        self.assertFalse(any('保留这些列' in warning for warning in sheet['warnings']))

    def test_sheet_column_and_total_business_row_limits(self):
        cases = [
            ('columns.csv', csv_bytes([BASIC + [''] * 92])),
            ('rows.csv', csv_bytes([BASIC] + [['', '', '', '项目', '', '', '', '', '']] * 5001)),
            ('sheets.xlsx', xlsx_bytes([(f'S{i}', [REPORT]) for i in range(21)])),
        ]
        for name, data in cases:
            with self.subTest(name=name), self.assertRaises(ValueError):
                parser.read_import(data, name)


class ProjectValidationTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(parser, '导入解析模块尚未实现')

    def test_normalization_keeps_only_supported_fields(self):
        project = parser.validate_project({'projectName': ' 项目 ', 'stage': 'EVT',
                                           'estimatedTotalDemand': '0', 'id': 99, 'unexpected': 'drop'})
        self.assertEqual(project['projectName'], '项目')
        self.assertEqual(project['estimatedTotalDemand'], 0)
        self.assertIsNone(project['expectedMpDate'])
        self.assertNotIn('id', project)
        self.assertNotIn('unexpected', project)
        self.assertEqual(len(project), 10)

    def test_text_limits_match_java_and_browser_utf16_counting(self):
        with self.assertRaises(ValueError):
            parser.validate_project({'projectName': '项目', 'application': '\U0001f600' * 128})
        with self.assertRaises(ValueError):
            parser.validate_project({'projectName': '项目', 'reportData': json.dumps(
                {'fields': {'备注': '\U0001f600' * 32000}}, ensure_ascii=False)})

    def test_rejects_missing_names_invalid_quantities_stages_dates_and_long_text(self):
        cases = [{}, {'projectName': '/'}, {'projectName': '项目', 'stage': 'MP'},
                 {'projectName': '项目', 'testDemandQuantity': True},
                 {'projectName': '项目', 'estimatedTotalDemand': -1},
                 {'projectName': '项目', 'estimatedTotalDemand': 9007199254740992},
                 {'projectName': '项目', 'estimatedTotalDemand': 1.5},
                 {'projectName': '项目', 'expectedMpDate': '2026年Q4'},
                 {'projectName': '项目', 'expectedMpDate': '2026-02-30'},
                 {'projectName': '项目', 'application': 'a' * 256},
                 {'projectName': '项目', 'reportData': 'invalid json'},
                 {'projectName': '项目', 'reportData': json.dumps({'fields': {'备注': 'a' * 64000}})}]
        for project in cases:
            with self.subTest(project_keys=list(project)), self.assertRaises(ValueError):
                parser.validate_project(project)


if __name__ == '__main__':
    unittest.main()
