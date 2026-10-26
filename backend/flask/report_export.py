"""Export the same normalized report rows used by the app, as a new workbook."""
import io
import re
from decimal import Decimal, InvalidOperation
from datetime import datetime, timezone, timedelta
import requests
from flask import Response, jsonify
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

KEYS = ['owner','reportDate','productLine','caseStatus','customerName','terminalCustomer','projectName','application','materialModel','specification','quotePrice','estimatedAnnualQuantity','estimatedAnnualRevenue','statusNote','expectedMpText','competitor','competitorPrice','notes']
LABELS = ['负责人','日期','产品线','案子状况','客户','终端客户','项目名称','应用','材料型号','成品规格','单价','预估年用量','预估年营业额','项目状态','预估量产时间','竞争对手','竞争对手单价','备注']


def build_report_workbook(rows):
    if not isinstance(rows,list): raise ValueError('报告数据格式异常，未导出不完整的文件')
    book=Workbook();book.remove(book.active)
    widths=[13,16,14,12,22,22,28,22,24,28,18,18,20,48,22,22,20,36]
    line=Side(style='thin',color='B3CCD8')
    for name in ['开发','成交','EOL','重点项目']:
        sheet=book.create_sheet(name);sheet.merge_cells('A1:R1')
        sheet['A1']=f'热管理事业部项目开发报告 · {name}';sheet['A1'].font=Font(name='微软雅黑',size=16,bold=True,color='FFFFFF');sheet['A1'].fill=PatternFill('solid',fgColor='145477');sheet['A1'].alignment=Alignment(horizontal='center',vertical='center');sheet.row_dimensions[1].height=34
        sheet.append(LABELS);sheet.row_dimensions[2].height=34
        for cell,width in zip(sheet[2],widths):
            cell.font=Font(name='微软雅黑',size=11,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='276C8C');cell.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center');cell.border=Border(left=line,right=line,top=line,bottom=line)
            sheet.column_dimensions[cell.column_letter].width=width
        sheet.freeze_panes='H3';sheet.print_title_rows='1:2'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
        sheet.oddFooter.center.text='第 &P 页 / 共 &N 页';sheet.sheet_view.showGridLines=False
    for item in rows:
        if not isinstance(item,dict) or item.get('category') not in ['开发','成交','EOL'] or not isinstance(item.get('fields'),dict): raise ValueError('报告分类或行格式异常，请先检查再导出')
        targets=[item['category']]+(['重点项目'] if item.get('important') else [])
        for name in targets:
            sheet=book[name];r=sheet.max_row+1
            for c,key in enumerate(KEYS,1):
                value=str(item['fields'].get(key) or '')
                if len(value)>32767 or ILLEGAL_CHARACTERS_RE.search(value): raise ValueError(f'第{r}行“{LABELS[c-1]}”包含Excel无法完整保存的内容，请先检查')
                cell=sheet.cell(r,c);cell.value=value;cell.data_type='s'
                if c in [11,12,13,17] and re.fullmatch(r'\d+(?:\.\d+)?',value):
                    try:
                        number=Decimal(value)
                        places=max(0,-number.as_tuple().exponent)
                        if len(number.as_tuple().digits)<=15 and places<=6:
                            cell.value=number;cell.number_format='#,##0'+('.'+'0'*places if places else '')
                    except InvalidOperation: pass
                cell.font=Font(name='微软雅黑',size=10,color='173F54');cell.alignment=Alignment(vertical='top',wrap_text=True);cell.fill=PatternFill('solid',fgColor='EDF6FA' if r%2 else 'FFFFFF');cell.border=Border(left=line,right=line,top=line,bottom=line)
                if key in ['quotePrice','estimatedAnnualRevenue','competitorPrice']:
                    cell.comment=Comment('币种：'+str(item.get('currency') or '原资料未提供')+'\n数量单位：'+str(item.get('quantityUnit') or '原资料未提供')+'\n计价说明：'+str(item.get('priceTerms') or '按原值保留，未换算或推算')+'\n报价及预估营业额不代表实际订单或收款。','项目业务工作台')
            sheet.row_dimensions[r].height=52
    for sheet in book:
        sheet.auto_filter.ref=f'A2:R{max(2,sheet.max_row)}';sheet.print_area=f'A1:R{max(2,sheet.max_row)}'
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


def register_report_export(app,spring_api):
    @app.get('/api/export/report.xlsx')
    def export_report_workbook():
        try:
            response=requests.get(f'{spring_api}/report/rows',timeout=30);response.raise_for_status()
            content=build_report_workbook(response.json())
        except requests.RequestException:
            return jsonify(error='无法读取项目报告数据，请检查主数据服务是否启动'),502
        except (ValueError,TypeError) as exc:
            return jsonify(error=str(exc)),400
        stamp=datetime.now(timezone(timedelta(hours=8))).strftime('%Y%m%d')
        return Response(content,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename=project-development-report-{stamp}.xlsx'})
